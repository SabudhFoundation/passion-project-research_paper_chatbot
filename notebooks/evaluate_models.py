'''
python notebooks/evaluate_models.py
'''

from __future__ import annotations

import sys
import time
from pathlib import Path
from typing import Dict, Tuple

import pandas as pd
from openpyxl import load_workbook

# PATH SETUP

REPO_ROOT = Path(__file__).resolve().parent.parent
SRC_PATH = REPO_ROOT / "src"

if str(SRC_PATH) not in sys.path:
    sys.path.insert(0, str(SRC_PATH))

if str(REPO_ROOT) not in sys.path:
    sys.path.insert(0, str(REPO_ROOT))

# PROJECT IMPORTS

from src.main import CONVERSATION_MEMORY, MainConfig, run_pipeline
from src.config import (
    EMBED_MODEL,
    MODEL_OPTIONS,
    TEMPERATURE,
    TOP_K,
    VECTOR_STORE_PATH,
)

# PATHS

INPUT_EXCEL = REPO_ROOT / "data" / "evaluation" / "Evaluation_Dataset.xlsx"
OUTPUT_EXCEL = REPO_ROOT / "data" / "evaluation" / "evaluation_results.xlsx"
TEMP_JSON = REPO_ROOT / "data" / "evaluation" / "temp.json"

# PROVIDERS

PROVIDERS = {
    "Gemini": MODEL_OPTIONS["Gemini: gemini-2.5-flash"],
    "Groq": MODEL_OPTIONS["Groq: llama-3.3-70b-versatile"],
    "HuggingFace": ("huggingface", "Qwen/Qwen2.5-72B-Instruct"),
}

MODEL_COLUMNS = {
    "Gemini": "gemini-2.5-flash Answer",
    "Groq": "llama-3.3-70b-versatile Answer",
    "HuggingFace": "Qwen2.5-72B Answer",
}

# HELPERS


def find_actual_header_row(excel_path: Path) -> int:
    """Automatically detect dataset header row."""

    workbook = load_workbook(
        excel_path,
        read_only=True,
        data_only=True,
    )

    sheet = workbook.active

    for idx, row in enumerate(sheet.iter_rows(values_only=True), start=1):
        values = [str(v).strip() if v is not None else "" for v in row]

        if "Question" in values and "Ground Truth Answer" in values:
            workbook.close()
            return idx

    workbook.close()

    raise ValueError("Could not detect Excel header row.")


def load_dataset(excel_path: Path) -> pd.DataFrame:
    header_row = find_actual_header_row(excel_path)

    df = pd.read_excel(
        excel_path,
        header=header_row - 1,
    )

    df = df.dropna(how="all")

    if "Question" not in df.columns:
        raise ValueError("Column 'Question' not found.")

    df = df[df["Question"].notna()]

    return df.reset_index(drop=True)


def initialize_columns(df: pd.DataFrame) -> pd.DataFrame:
    for column in MODEL_COLUMNS.values():
        if column not in df.columns:
            df[column] = ""

    return df


def truncate_excel_text(text: str, limit: int = 30000) -> str:
    """Prevent oversized Excel cells."""

    text = str(text)

    if len(text) <= limit:
        return text

    return text[:limit] + "\n\n[TRUNCATED]"


def save_progress(df: pd.DataFrame) -> None:
    OUTPUT_EXCEL.parent.mkdir(parents=True, exist_ok=True)
    
    # Retry logic for handling locked files
    max_retries = 3
    for attempt in range(max_retries):
        try:
            df.to_excel(OUTPUT_EXCEL, index=False)
            return
        except PermissionError as e:
            if attempt < max_retries - 1:
                print(f"File locked. Retrying in 2 seconds... (Attempt {attempt + 1}/{max_retries})")
                time.sleep(2)
            else:
                print(f"ERROR: Could not write to {OUTPUT_EXCEL}")
                print("SOLUTION: Close the Excel file if it's open, then run the script again.")
                raise


def is_quota_error(error_message: str) -> bool:
    """Detect quota/rate-limit related errors."""

    text = str(error_message).lower()

    quota_patterns = [
        "quota exceeded",
        "rate limit exceeded",
        "too many requests",
        "resource_exhausted",
        "insufficient quota",
        "you exceeded your current quota",
        "daily limit exceeded",
        "usage limit exceeded",
        "429",
    ]

    return any(pattern in text for pattern in quota_patterns)


def should_rerun(value) -> bool:
    """Decide whether answer should be regenerated."""

    if pd.isna(value):
        return True

    text = str(value).strip()

    if not text:
        return True

    lowered = text.lower()

    error_patterns = [
        "error",
        "quota",
        "rate limit",
        "too many requests",
        "resource_exhausted",
        "429",
    ]

    return any(pattern in lowered for pattern in error_patterns)


def ask_model(
    question: str,
    provider: str,
    model_name: str,
) -> Tuple[str, bool]:
    """
    Returns:
        (response_text, is_quota_error)
    """

    CONVERSATION_MEMORY.clear()

    vector_store_path = Path(VECTOR_STORE_PATH)

    if not vector_store_path.is_absolute():
        vector_store_path = REPO_ROOT / vector_store_path

    max_attempts = 3
    backoff = 2

    for attempt in range(1, max_attempts + 1):

        config = MainConfig(
            question=question,
            provider=provider,
            model_name=model_name,
            temperature=TEMPERATURE,
            vector_store_path=vector_store_path,
            output_json_path=TEMP_JSON,
            top_k=TOP_K,
            embed_model=EMBED_MODEL,
        )

        try:
            result = run_pipeline(config)

            if result.get("error"):
                error_message = str(result["error"])

                if is_quota_error(error_message):
                    return (f"QUOTA_ERROR: {error_message}", True)

                return (f"ERROR: {error_message}", False)

            answer = str(result.get("answer", "")).strip()

            if not answer:
                raise RuntimeError("Empty response generated.")

            return (answer, False)

        except Exception as exc:
            error_message = str(exc)

            retryable = any(
                keyword in error_message.lower()
                for keyword in [
                    "timeout",
                    "connection",
                    "temporarily",
                    "server error",
                    "502",
                    "503",
                    "504",
                ]
            )

            if retryable and attempt < max_attempts:
                print(
                    f"[Retry {attempt}/{max_attempts}] "
                    f"{provider} -> {error_message}"
                )

                time.sleep(backoff)
                backoff *= 2
                continue

            if is_quota_error(error_message):
                return (f"QUOTA_ERROR: {error_message}", True)

            return (f"ERROR: {error_message}", False)

    return ("ERROR: Unknown failure", False)


def evaluate_provider(
    df: pd.DataFrame,
    row_index: int,
    question: str,
    provider_key: str,
    disabled_providers: Dict[str, str],
) -> None:

    column_name = MODEL_COLUMNS[provider_key]

    # Skip provider after quota exhaustion
    if provider_key in disabled_providers:
        quota_message = disabled_providers[provider_key]

        print(f"[{provider_key}] SKIPPED -> quota disabled")

        df.at[row_index, column_name] = quota_message
        return

    provider, model_name = PROVIDERS[provider_key]

    print(f"[{provider_key}] Processing...")

    answer, is_quota = ask_model(
        question=question,
        provider=provider,
        model_name=model_name,
    )

    if is_quota:
        disabled_providers[provider_key] = answer

        print(f"[{provider_key}] QUOTA EXCEEDED")

        df.at[row_index, column_name] = answer

        # Persist immediately
        save_progress(df)

        print(f"[{provider_key}] Quota state saved")
        return

    if answer.startswith("ERROR"):
        print(f"[{provider_key}] FAILED")

        df.at[row_index, column_name] = truncate_excel_text(answer)
        return

    df.at[row_index, column_name] = answer

    print(f"[{provider_key}] DONE")


# MAIN


def generate_evaluation_sheet() -> None:

    print("\nLoading dataset...")

    if OUTPUT_EXCEL.exists():

        print("\nResume file detected.")

        print(
            f"Loading existing progress from:\n"
            f"{OUTPUT_EXCEL}"
        )

        df = pd.read_excel(OUTPUT_EXCEL)

    else:

        print("\nNo previous progress found.")

        print("Starting fresh evaluation.")

        df = load_dataset(INPUT_EXCEL)

        df = initialize_columns(df)

        save_progress(df)

    # Ensure columns exist even in old files
    df = initialize_columns(df)

    total_questions = len(df)

    print(f"\nLoaded {total_questions} questions")

    disabled_providers = {}

    for idx, row in df.iterrows():

        question = str(row["Question"]).strip()

        if not question:
            continue

        print("\n" + "=" * 100)
        print(f"[QUESTION {idx + 1}/{total_questions}]")
        print(question)
        print("=" * 100)

        for provider_key in PROVIDERS.keys():

            column_name = MODEL_COLUMNS[provider_key]

            current_value = row.get(column_name, "")

            if not should_rerun(current_value):

                print(
                    f"[{provider_key}] SKIPPED -> already completed"
                )

                continue

            print(f"[{provider_key}] NEEDS RERUN")

            evaluate_provider(
                df=df,
                row_index=idx,
                question=question,
                provider_key=provider_key,
                disabled_providers=disabled_providers,
            )

            time.sleep(0.3)

        save_progress(df)

        print("\nProgress saved")

    print("\nEvaluation completed")
    print(f"Saved to: {OUTPUT_EXCEL}")


# ENTRY

if __name__ == "__main__":
    generate_evaluation_sheet()