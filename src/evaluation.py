"""
evaluation.py

Evaluate saved RAG answers using DeepEval and export results to Excel.

This file is import-safe:
- Importing evaluation.py will NOT run evaluation.
- Evaluation runs only when executed directly.

Example:
python src/evaluation.py --input_json data/results/output.json --output_excel data/results/evaluation_results.xlsx
"""

from __future__ import annotations

import argparse
import json
import os
from pathlib import Path
from typing import Any

import pandas as pd
from dotenv import load_dotenv
from langchain_groq import ChatGroq
from openpyxl import load_workbook
from openpyxl.styles import Alignment

from deepeval.metrics import (
    AnswerRelevancyMetric,
    ContextualPrecisionMetric,
    ContextualRecallMetric,
    ContextualRelevancyMetric,
    FaithfulnessMetric,
    HallucinationMetric,
)
from deepeval.models import DeepEvalBaseLLM
from deepeval.test_case import LLMTestCase


load_dotenv()


# =========================================================
# Custom Groq Evaluator
# =========================================================

class GroqEvaluator(DeepEvalBaseLLM):
    def __init__(
        self,
        model_name: str = "llama-3.3-70b-versatile",
        temperature: float = 0.0,
    ) -> None:
        groq_api_key = os.getenv("GROQ_API_KEY")

        if not groq_api_key:
            raise EnvironmentError(
                "GROQ_API_KEY is not set in your environment or .env file."
            )

        self.model_name = model_name
        self.model = ChatGroq(
            groq_api_key=groq_api_key,
            model_name=model_name,
            temperature=temperature,
        )

    def load_model(self):
        return self.model

    def generate(self, prompt: str) -> str:
        return self.model.invoke(prompt).content

    async def a_generate(self, prompt: str) -> str:
        return self.model.invoke(prompt).content

    def get_model_name(self) -> str:
        return self.model_name


# =========================================================
# Config
# =========================================================

class EvaluationConfig:
    def __init__(
        self,
        input_json: Path,
        output_excel: Path,
        evaluator_model: str,
    ) -> None:
        self.input_json = input_json
        self.output_excel = output_excel
        self.evaluator_model = evaluator_model


def parse_args() -> EvaluationConfig:
    parser = argparse.ArgumentParser(
        description="Evaluate RAG outputs using DeepEval.",
    )

    parser.add_argument(
        "--input_json",
        type=str,
        default="data/results/output.json",
        help="Path to saved RAG output JSON.",
    )

    parser.add_argument(
        "--output_excel",
        type=str,
        default="data/results/evaluation_results.xlsx",
        help="Path where evaluation Excel file will be saved.",
    )

    parser.add_argument(
        "--evaluator_model",
        type=str,
        default="llama-3.3-70b-versatile",
        help="Groq model used as the DeepEval evaluator.",
    )

    args = parser.parse_args()

    return EvaluationConfig(
        input_json=Path(args.input_json),
        output_excel=Path(args.output_excel),
        evaluator_model=args.evaluator_model,
    )


# =========================================================
# Data helpers
# =========================================================

def load_output_json(path: Path) -> list[dict[str, Any]]:
    if not path.exists():
        raise FileNotFoundError(f"Output JSON not found: {path}")

    with path.open("r", encoding="utf-8") as file:
        data = json.load(file)

    if isinstance(data, dict):
        data = [data]

    if not isinstance(data, list):
        raise ValueError("Output JSON must contain either a dict or a list of dicts.")

    return [
        item
        for item in data
        if isinstance(item, dict)
    ]


def extract_context_list(item: dict[str, Any]) -> list[str]:
    retrieved_chunks = item.get("retrieved_chunks", [])

    if not isinstance(retrieved_chunks, list):
        return []

    context_list: list[str] = []

    for chunk in retrieved_chunks:
        if not isinstance(chunk, dict):
            continue

        content = (
            chunk.get("content")
            or chunk.get("text")
            or chunk.get("embedding_text")
            or ""
        )

        content = str(content).strip()

        if content:
            context_list.append(content)

    return context_list


def safe_score(metric: Any) -> float | None:
    score = getattr(metric, "score", None)

    if score is None:
        return None

    try:
        return round(float(score), 3)
    except Exception:
        return None


def safe_success(metric: Any) -> bool | None:
    try:
        return bool(metric.is_successful())
    except Exception:
        return None


def safe_reason(metric: Any) -> str:
    return str(getattr(metric, "reason", "") or "")


# =========================================================
# Metric setup
# =========================================================

def build_metrics(evaluator_llm: GroqEvaluator) -> dict[str, Any]:
    return {
        "answer_relevancy": AnswerRelevancyMetric(
            threshold=0.7,
            model=evaluator_llm,
        ),
        "faithfulness": FaithfulnessMetric(
            threshold=0.7,
            model=evaluator_llm,
        ),
        "hallucination": HallucinationMetric(
            threshold=0.5,
            model=evaluator_llm,
        ),
        "context_precision": ContextualPrecisionMetric(
            threshold=0.7,
            model=evaluator_llm,
        ),
        "context_recall": ContextualRecallMetric(
            threshold=0.7,
            model=evaluator_llm,
        ),
        "context_relevancy": ContextualRelevancyMetric(
            threshold=0.7,
            model=evaluator_llm,
        ),
    }


# =========================================================
# Evaluation
# =========================================================

def evaluate_item(
    item: dict[str, Any],
    metrics: dict[str, Any],
) -> dict[str, Any]:
    question = str(item.get("question", "") or "").strip()
    answer = str(item.get("answer", "") or "").strip()
    route = str(item.get("route", "") or "").strip()
    context_list = extract_context_list(item)

    if not question or not answer:
        return {
            "Question": question,
            "Answer": answer,
            "Route": route,
            "Skipped": True,
            "Skip Reason": "Missing question or answer.",
        }

    if not context_list:
        context_list = ["No retrieved context was saved for this item."]

    print("\n===================================")
    print("QUESTION:", question)
    print("===================================")

    test_case = LLMTestCase(
        input=question,
        actual_output=answer,
        expected_output=answer,
        context=context_list,
        retrieval_context=context_list,
    )

    for metric in metrics.values():
        metric.measure(test_case)

    answer_metric = metrics["answer_relevancy"]
    faithfulness_metric = metrics["faithfulness"]
    hallucination_metric = metrics["hallucination"]
    context_precision_metric = metrics["context_precision"]
    context_recall_metric = metrics["context_recall"]
    context_relevancy_metric = metrics["context_relevancy"]

    return {
        "Question": question,
        "Answer": answer,
        "Route": route,
        "Context": "\n\n-------------------------\n\n".join(context_list),
        "Skipped": False,
        "Skip Reason": "",

        "Answer Relevancy Score": safe_score(answer_metric),
        "Answer Relevancy Passed": safe_success(answer_metric),
        "Answer Relevancy Reason": safe_reason(answer_metric),

        "Faithfulness Score": safe_score(faithfulness_metric),
        "Faithfulness Passed": safe_success(faithfulness_metric),
        "Faithfulness Reason": safe_reason(faithfulness_metric),

        "Hallucination Score": safe_score(hallucination_metric),
        "Hallucination Passed": safe_success(hallucination_metric),
        "Hallucination Reason": safe_reason(hallucination_metric),

        "Context Precision Score": safe_score(context_precision_metric),
        "Context Precision Passed": safe_success(context_precision_metric),
        "Context Precision Reason": safe_reason(context_precision_metric),

        "Context Recall Score": safe_score(context_recall_metric),
        "Context Recall Passed": safe_success(context_recall_metric),
        "Context Recall Reason": safe_reason(context_recall_metric),

        "Context Relevancy Score": safe_score(context_relevancy_metric),
        "Context Relevancy Passed": safe_success(context_relevancy_metric),
        "Context Relevancy Reason": safe_reason(context_relevancy_metric),
    }


def save_excel_report(
    rows: list[dict[str, Any]],
    output_excel: Path,
) -> Path:
    output_excel.parent.mkdir(parents=True, exist_ok=True)

    df = pd.DataFrame(rows)
    df.to_excel(output_excel, index=False)

    workbook = load_workbook(output_excel)
    worksheet = workbook.active

    if worksheet is None:
        workbook.save(output_excel)
        return output_excel

    for row in worksheet.iter_rows():
        for cell in row:
            cell.alignment = Alignment(
                wrap_text=True,
                vertical="top",
            )

    for column_cells in worksheet.columns:
        max_length = 0
        column_letter = column_cells[0].column_letter

        for cell in column_cells:
            try:
                if cell.value:
                    max_length = max(
                        max_length,
                        len(str(cell.value)),
                    )
            except Exception:
                pass

        worksheet.column_dimensions[column_letter].width = min(max_length + 5, 50)

    for row in worksheet.iter_rows():
        worksheet.row_dimensions[row[0].row].height = 80

    workbook.save(output_excel)
    return output_excel


def run_evaluation(config: EvaluationConfig) -> Path:
    data = load_output_json(config.input_json)

    if not data:
        raise ValueError(f"No valid evaluation rows found in {config.input_json}")

    evaluator_llm = GroqEvaluator(
        model_name=config.evaluator_model,
        temperature=0.0,
    )

    metrics = build_metrics(evaluator_llm)

    results: list[dict[str, Any]] = []

    for item in data:
        try:
            row = evaluate_item(
                item,
                metrics,
            )
        except Exception as exc:
            row = {
                "Question": item.get("question", ""),
                "Answer": item.get("answer", ""),
                "Route": item.get("route", ""),
                "Skipped": True,
                "Skip Reason": f"Evaluation failed: {exc}",
            }

        results.append(row)

    return save_excel_report(
        rows=results,
        output_excel=config.output_excel,
    )


def main() -> None:
    config = parse_args()

    print("\n===================================")
    print("Evaluation Starting")
    print("Input JSON:", config.input_json)
    print("Output Excel:", config.output_excel)
    print("Evaluator Model:", config.evaluator_model)
    print("===================================")

    output_path = run_evaluation(config)

    print("\n===================================")
    print("Evaluation Completed Successfully")
    print("Excel Saved At:")
    print(output_path)
    print("===================================")


if __name__ == "__main__":
    main()